"""
ADAPT — DOCX Document Parser

Extracts text, headings, lists, and tables from DOCX files using python-docx.
"""

from pathlib import Path
from typing import Any, Dict, List
import docx

from app.core.exceptions import FileParsingError
from app.core.logging import get_logger

logger = get_logger("rag.docx_parser")


class DOCXParser:
    """Extracts text, structured paragraphs and tables from DOCX files."""

    @staticmethod
    def parse(file_path: str) -> Dict[str, Any]:
        path = Path(file_path)
        if not path.exists():
            raise FileParsingError(f"DOCX file not found at {file_path}")

        try:
            doc = docx.Document(file_path)
            paragraphs_text: List[str] = []
            structured_blocks: List[Dict[str, Any]] = []

            for p in doc.paragraphs:
                text = p.text.strip()
                if not text:
                    continue
                paragraphs_text.append(text)
                structured_blocks.append({
                    "type": "paragraph",
                    "style": p.style.name if p.style else "Normal",
                    "text": text,
                })

            # Also parse tables
            for t_idx, table in enumerate(doc.tables):
                table_rows = []
                for row in table.rows:
                    row_cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                    if any(row_cells):
                        table_rows.append(" | ".join(row_cells))
                if table_rows:
                    table_str = "\n".join(table_rows)
                    paragraphs_text.append(table_str)
                    structured_blocks.append({
                        "type": "table",
                        "index": t_idx,
                        "text": table_str,
                    })

            full_text = "\n\n".join(paragraphs_text)

            return {
                "full_text": full_text,
                "blocks": structured_blocks,
                "title": path.stem,
                "page_count": max(1, len(paragraphs_text) // 25), # approximate
            }

        except Exception as e:
            logger.error(f"Failed to parse DOCX {file_path}: {e}")
            raise FileParsingError(f"Could not parse DOCX: {str(e)}")
